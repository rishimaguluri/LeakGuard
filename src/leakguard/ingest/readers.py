"""CSV and XLSX readers.

Everything is read as text (Excel dates become ISO strings) so the mapping
step sees one consistent format. Handles unknown delimiters and encodings,
multiple sheets, and title rows above the header.
"""

from __future__ import annotations

import csv
import io
import re
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from charset_normalizer import from_bytes
from openpyxl import load_workbook

SUPPORTED_EXTENSIONS = {".csv", ".txt", ".tsv", ".xlsx", ".xlsm"}
HEADER_SCAN_ROWS = 25


class UnsupportedFile(Exception):
    """The file type cannot be read in v1."""


@dataclass
class RawTable:
    header: list[str]
    rows: list[list[str]]
    first_row_number: int  # 1-based file row number of rows[0]
    sheet: str | None = None
    encoding: str | None = None
    delimiter: str | None = None
    sheet_names: list[str] | None = None


def header_key(text: str) -> str:
    """Case and punctuation insensitive key used to compare headers."""
    return re.sub(r"[^a-z0-9]", "", str(text).lower())


# Low level grid readers ---------------------------------------------------------


def decode_bytes(data: bytes) -> tuple[str, str]:
    """Return (text, encoding).

    Order: UTF-16 with a byte order mark, UTF-8, then Windows-1252 (what
    Excel on Windows writes for most US and European exports), then
    automatic detection. Detection goes last because it guesses badly on
    short files.
    """
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16"), "utf-16"
    for encoding, label in (("utf-8-sig", "utf-8"), ("cp1252", "cp1252")):
        try:
            return data.decode(encoding), label
        except UnicodeDecodeError:
            pass
    best = from_bytes(data).best()
    if best is not None and best.encoding:
        try:
            return data.decode(best.encoding), best.encoding
        except (UnicodeDecodeError, LookupError):
            pass
    return data.decode("cp1252", errors="replace"), "cp1252"


def detect_delimiter(sample: str) -> str:
    """Pick the delimiter that splits the most lines into the same number
    of columns. (csv.Sniffer can hang on some inputs, so it is not used.)"""
    lines = [line for line in sample.splitlines() if line.strip()][:40]
    best, best_score = ",", (-1, -1)
    for delimiter in ",;\t|":
        widths = [len(row) for row in csv.reader(lines, delimiter=delimiter)]
        if not widths:
            continue
        width, freq = Counter(widths).most_common(1)[0]
        if width < 2:
            continue
        if (freq, width) > best_score:
            best, best_score = delimiter, (freq, width)
    return best


def read_csv_grid(path: Path) -> tuple[list[list[str]], str, str]:
    text, encoding = decode_bytes(path.read_bytes())
    sample = "\n".join(text.splitlines()[:50])
    delimiter = detect_delimiter(sample)
    grid = [list(row) for row in csv.reader(io.StringIO(text), delimiter=delimiter)]
    return grid, encoding, delimiter


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.hour == value.minute == value.second == 0:
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        if value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        return repr(value)
    return str(value)


def list_sheets(path: Path) -> list[str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def read_xlsx_grid(path: Path, sheet: str | None = None) -> tuple[list[list[str]], str, list[str]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        names = list(wb.sheetnames)
        if sheet and sheet not in names:
            raise ValueError(f"Sheet '{sheet}' not found. Sheets in this file: {', '.join(names)}")
        ws = wb[sheet] if sheet else wb[names[0]]
        grid = [[_cell_text(v) for v in row] for row in ws.iter_rows(values_only=True)]
        return grid, ws.title, names
    finally:
        wb.close()


# Header detection -----------------------------------------------------------------


def _looks_like_header(row: list[str]) -> bool:
    cells = [c.strip() for c in row if c and c.strip()]
    if len(cells) < 2:
        return False
    texty = sum(
        1 for c in cells if re.search(r"[A-Za-z]", c) and not re.fullmatch(r"[\d$.,\-/ ]+", c)
    )
    return texty / len(cells) >= 0.8


def detect_header_row(grid: list[list[str]], known_headers: set[str] | None = None) -> int:
    """Index of the header row.

    With known header names (from a mapping), picks the row matching most of
    them. Without, picks the first row that looks like column names.
    """
    scan = grid[:HEADER_SCAN_ROWS]
    if known_headers:
        best_index, best_score = -1, 0
        for i, row in enumerate(scan):
            score = sum(1 for c in row if header_key(c) in known_headers)
            if score > best_score:
                best_index, best_score = i, score
        if best_score >= 2:
            return best_index
    for i, row in enumerate(scan):
        if _looks_like_header(row):
            return i
    return 0


# Public API ---------------------------------------------------------------------


def read_table(
    path: Path,
    known_headers: set[str] | None = None,
    sheet: str | None = None,
    header_row: int | str = "auto",
) -> RawTable:
    """Read a CSV or XLSX file into a header plus rows of text."""
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        if ext == ".pdf":
            raise UnsupportedFile(
                "PDF files are not supported in v1. Ask the hotel for the same report as CSV or Excel."
            )
        if ext == ".xls":
            raise UnsupportedFile(
                "Old .xls files are not supported. Open in Excel and save as .xlsx."
            )
        raise UnsupportedFile(f"'{ext}' files are not supported. Use CSV or XLSX.")

    encoding = delimiter = sheet_used = None
    sheet_names = None
    if ext in (".xlsx", ".xlsm"):
        grid, sheet_used, sheet_names = read_xlsx_grid(path, sheet)
    else:
        grid, encoding, delimiter = read_csv_grid(path)

    if not grid:
        return RawTable([], [], 1, sheet_used, encoding, delimiter, sheet_names)
    if isinstance(header_row, int):
        index = max(header_row - 1, 0)
    else:
        index = detect_header_row(grid, known_headers)
    header = [c.strip() for c in grid[index]]
    width = len(header)
    rows = [(r + [""] * width)[:width] for r in grid[index + 1 :]]
    return RawTable(header, rows, index + 2, sheet_used, encoding, delimiter, sheet_names)
