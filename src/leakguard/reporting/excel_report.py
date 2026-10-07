"""Excel audit workbook: Summary, By property, By management company,
Line items, Method. Built from ReportData so it matches the dashboard."""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from leakguard.reporting import metrics as M
from leakguard.reporting.report_data import ReportData
from leakguard.schemas import STATUS_LABELS

NAVY = "1A3766"
TEXT = "0E2240"
TEXT_2 = "25406A"
LINE = "BFD0E6"
CARD = "F8FAFD"
HEAD_FILL = "E1E9F4"
LOSS = "8C2B45"
GAIN = "1F5A3D"

USD = '"$"#,##0'
USD_CENTS = '"$"#,##0.00'
PCT = "0.0%"
DATE = "mmm d, yyyy"

TITLE = Font(name="Georgia", size=18, color=TEXT)
H2 = Font(name="Georgia", size=13, color=TEXT)
HEAD = Font(name="Calibri", size=11, bold=True, color=TEXT)
BODY = Font(name="Calibri", size=11, color=TEXT)
MUTED = Font(name="Calibri", size=10, color=TEXT_2, italic=True)
THIN = Border(bottom=Side(style="thin", color=LINE))


def clean(value: object) -> object:
    """Excel cannot store NaN; blank the cell instead. numpy scalars become plain Python."""
    if value is None:
        return None
    if isinstance(value, float) and value != value:
        return None
    if hasattr(value, "item") and not isinstance(value, str):
        value = value.item()
        if isinstance(value, float) and value != value:
            return None
    return value


def _header(ws: Worksheet, row: int, labels: list[str]) -> None:
    for col, label in enumerate(labels, start=1):
        cell = ws.cell(row=row, column=col, value=label)
        cell.font = HEAD
        cell.fill = PatternFill("solid", fgColor=HEAD_FILL)
        cell.border = THIN
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _widths(ws: Worksheet, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _put(
    ws: Worksheet, row: int, col: int, value: object, fmt: str | None = None, font: Font = BODY
) -> None:
    cell = ws.cell(row=row, column=col, value=clean(value))
    cell.font = font
    if fmt:
        cell.number_format = fmt


def _summary(ws: Worksheet, r: ReportData) -> None:
    s, n = r.summary, r.noi
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    _widths(ws, [44, 20, 16, 16])
    _put(ws, 1, 1, "LeakGuard virtual card audit", font=TITLE)
    _put(ws, 2, 1, f"{r.portfolio_name} ({r.owner_name})", font=H2)
    _put(
        ws,
        3,
        1,
        f"Period: {M.fmt_date(r.start)} to {M.fmt_date(r.end)}. Data as of {M.fmt_date(r.as_of)}.",
        font=BODY,
    )
    if r.data_mode == "demo":
        _put(ws, 4, 1, "Demo data. Fictional portfolio for illustration.", font=MUTED)

    rows = [
        ("Virtual cards reviewed", s.vcc_count, "#,##0"),
        ("OTA-collect revenue reviewed", s.vcc_volume_cents / 100, USD),
        ("Dollars leaked", s.leaked_cents / 100, USD),
        ("Leak rate", s.leak_rate, PCT),
        ("Recovered to date", s.recovered_cents / 100, USD),
        ("Open dollars at risk", s.open_at_risk_cents / 100, USD),
        ("Written off", s.written_off_cents / 100, USD),
        ("Recoverable estimate (estimate)", s.recoverable_estimate_cents / 100, USD),
        (f"LeakGuard fee at {r.fee_pct * 100:.0f}% of recovered cash", s.fee_cents / 100, USD),
        ("Owner net from recovered cash", s.owner_net_cents / 100, USD),
        ("NOI impact, annualized owner net", n.annualized_owner_net_cents / 100, USD),
        (
            f"Implied value at {n.cap_rate * 100:.1f}% cap rate (estimate)",
            n.implied_value_cents / 100,
            USD,
        ),
    ]
    _put(ws, 6, 1, "Key figures", font=H2)
    _header(ws, 7, ["Measure", "Value"])
    for i, (label, value, fmt) in enumerate(rows, start=8):
        _put(ws, i, 1, label)
        _put(ws, i, 2, value, fmt)
        ws.cell(row=i, column=1).border = THIN
        ws.cell(row=i, column=2).border = THIN
    ws.cell(row=10, column=2).font = Font(name="Calibri", size=11, bold=True, color=LOSS)
    ws.cell(row=12, column=2).font = Font(name="Calibri", size=11, bold=True, color=GAIN)

    row = 8 + len(rows) + 1
    _put(ws, row, 1, "Dollars leaked by type", font=H2)
    _header(ws, row + 1, ["Exception type", "Leaked", "Items", "Recovered"])
    for i, t in enumerate(r.by_type.itertuples(), start=row + 2):
        _put(ws, i, 1, t.label)
        _put(ws, i, 2, t.leaked_cents / 100, USD)
        _put(ws, i, 3, int(t.count), "#,##0")
        _put(ws, i, 4, t.recovered_cents / 100, USD)
    row = row + 2 + len(r.by_type) + 1
    if r.takeaway:
        _put(ws, row, 1, r.takeaway, font=Font(name="Calibri", size=11, bold=True, color=TEXT))
        row += 1
    _put(
        ws,
        row + 1,
        1,
        "Estimates are labeled as estimates. Recovery is not guaranteed. "
        "The fee applies only to cash actually recovered.",
        font=MUTED,
    )


def _by_property(ws: Worksheet, r: ReportData) -> None:
    ws.sheet_view.showGridLines = False
    labels = [
        "Property",
        "Name",
        "Management company",
        "Rooms",
        "OTA-collect revenue",
        "Leaked",
        "Leak rate",
        "Recovered",
        "Open at risk",
        "Recovery rate",
        "Items",
    ]
    _widths(ws, [10, 34, 26, 8, 20, 14, 10, 14, 14, 13, 8])
    _header(ws, 1, labels)
    for i, p in enumerate(r.by_property.itertuples(), start=2):
        values = [
            p.property_code,
            p.property_name,
            p.company,
            int(p.room_count or 0),
            p.vcc_volume_cents / 100,
            p.leaked_cents / 100,
            p.leak_rate,
            p.recovered_cents / 100,
            p.open_cents / 100,
            p.recovery_rate,
            int(p.exceptions),
        ]
        fmts = [None, None, None, "#,##0", USD, USD, PCT, USD, USD, PCT, "#,##0"]
        for col, (v, f) in enumerate(zip(values, fmts, strict=True), start=1):
            _put(ws, i, col, v, f)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(labels))}{len(r.by_property) + 1}"


def _by_company(ws: Worksheet, r: ReportData) -> None:
    ws.sheet_view.showGridLines = False
    labels = [
        "Rank",
        "Management company",
        "Properties",
        "OTA-collect revenue",
        "Leaked",
        "Leak rate",
        "Expired-card share",
        "Avg days to resolve",
        "Open at risk",
        "Recovered",
    ]
    _widths(ws, [7, 28, 11, 20, 14, 11, 17, 18, 14, 14])
    _header(ws, 1, labels)
    for i, c in enumerate(r.scorecard.itertuples(), start=2):
        days = (
            None
            if c.avg_days_to_resolve != c.avg_days_to_resolve
            else round(float(c.avg_days_to_resolve), 1)
        )
        values = [
            int(c.rank),
            c.company,
            int(c.properties),
            c.vcc_volume_cents / 100,
            c.leaked_cents / 100,
            c.leak_rate,
            c.expired_share,
            days,
            c.open_cents / 100,
            c.recovered_cents / 100,
        ]
        fmts = [None, None, "#,##0", USD, USD, PCT, PCT, "0.0", USD, USD]
        for col, (v, f) in enumerate(zip(values, fmts, strict=True), start=1):
            _put(ws, i, col, v, f)
    if r.takeaway:
        _put(ws, len(r.scorecard) + 3, 1, r.takeaway, font=HEAD)


def _line_items(ws: Worksheet, r: ReportData) -> None:
    labels = [
        "Priority",
        "Property",
        "Management company",
        "OTA",
        "OTA confirmation",
        "PMS confirmation",
        "Guest last name",
        "Arrival",
        "Departure",
        "Card last 4",
        "Card amount",
        "Charged",
        "Settled",
        "Amount at risk",
        "Card expiry",
        "Days to expiry",
        "Exception type",
        "Match method",
        "Confidence",
        "Why flagged",
        "Status",
        "Assigned to",
        "Recovered",
        "Recovered date",
    ]
    _widths(
        ws,
        [
            9,
            9,
            22,
            11,
            16,
            16,
            16,
            12,
            12,
            10,
            13,
            13,
            13,
            14,
            12,
            12,
            22,
            12,
            11,
            70,
            13,
            24,
            13,
            13,
        ],
    )
    _header(ws, 1, labels)
    items = r.line_items
    for i, e in enumerate(items.itertuples(), start=2):
        days = (
            None
            if e.days_to_expiry is None or e.days_to_expiry != e.days_to_expiry
            else int(e.days_to_expiry)
        )
        values = [
            round(float(e.priority_score)),
            e.property_code,
            e.company,
            "Booking.com" if e.ota == "booking" else "Expedia",
            e.ota_confirmation_no,
            e.pms_confirmation_no,
            e.guest_last_name,
            e.arrival_date,
            e.departure_date,
            e.card_last4,
            e.expected_cents / 100,
            e.charged_cents / 100,
            None
            if e.settled_cents is None or e.settled_cents != e.settled_cents
            else e.settled_cents / 100,
            e.amount_at_risk_cents / 100,
            e.expiry_date,
            days,
            e.type_label,
            e.match_method,
            e.confidence,
            e.rule,
            STATUS_LABELS.get(e.status, e.status),
            e.assigned_to,
            e.recovered_cents / 100,
            e.recovered_date,
        ]
        fmts = [
            "#,##0",
            None,
            None,
            None,
            "@",
            "@",
            None,
            DATE,
            DATE,
            "@",
            USD_CENTS,
            USD_CENTS,
            USD_CENTS,
            USD_CENTS,
            DATE,
            "0",
            None,
            None,
            None,
            None,
            None,
            None,
            USD_CENTS,
            DATE,
        ]
        for col, (v, f) in enumerate(zip(values, fmts, strict=True), start=1):
            ws.cell(row=i, column=col, value=clean(v)).number_format = f or "General"
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(labels))}{max(len(items) + 1, 2)}"


def _method(ws: Worksheet, r: ReportData) -> None:
    ws.sheet_view.showGridLines = False
    _widths(ws, [28, 110])
    _put(ws, 1, 1, "Method and limits", font=TITLE)
    for i, (title, body) in enumerate(r.method, start=3):
        _put(ws, i, 1, title, font=HEAD)
        cell = ws.cell(row=i, column=2, value=body)
        cell.font = BODY
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 62
        ws.cell(row=i, column=1).alignment = Alignment(vertical="top")


def build_workbook(r: ReportData) -> bytes:
    wb = Workbook()
    _summary(wb.active, r)
    _by_property(wb.create_sheet("By property"), r)
    _by_company(wb.create_sheet("By management company"), r)
    _line_items(wb.create_sheet("Line items"), r)
    _method(wb.create_sheet("Method"), r)
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = NAVY
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
