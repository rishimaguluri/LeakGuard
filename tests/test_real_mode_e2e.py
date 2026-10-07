"""End to end: real mode with new files and zero code changes.

1. Point LeakGuard at an empty temp project (config, mappings, data).
2. Add a real-mode portfolio to config/portfolios.yaml.
3. Drop a "fake real" hotel's exports into data/raw/<portfolio>/<property>/.
4. Convert its processor export into a format LeakGuard has never seen and
   add only a new YAML mapping for it.
5. Set data_mode: real in settings.yaml, then run import, match and report
   through the CLI, exactly as a person would.
"""

import csv
from datetime import date
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook
from sqlalchemy import select

from leakguard import __main__ as cli
from leakguard.config import MODE_ENV_VAR, load_settings
from leakguard.db import dispose_engine, session_scope
from leakguard.demo.generator import FAKE_REAL_PORTFOLIO, write_fake_real_hotel
from leakguard.models import Exception_, SourceFile

from .conftest import make_env

AS_OF = date(2026, 6, 30)

NEW_GATEWAY_MAPPING = {
    "name": "acme_gateway_settlement",
    "source_name": "processor_settlement",
    "target": "processor_transactions",
    "header_row": "auto",
    "columns": {
        "transaction_date": ["Txn Posted"],
        "settlement_date": ["Funded On"],
        "card_last4": ["Masked PAN"],
        "amount": ["Gross Amt"],
        "transaction_type": ["Kind"],
        "auth_code": ["Approval"],
        "reference": ["Folio Ref"],
    },
    "defaults": {"currency": "USD"},
    "notes": "Test mapping for a gateway format the shipped mappings do not cover.",
}
RENAMES = {
    "Transaction Date": "Txn Posted",
    "Settlement Date": "Funded On",
    "Card Number": "Masked PAN",
    "Amount": "Gross Amt",
    "Transaction Type": "Kind",
    "Auth Code": "Approval",
    "Reference": "Folio Ref",
}


def _convert_processor_to_new_format(folder: Path) -> None:
    """Rewrite the processor XLSX as a CSV with different headers and dates."""
    [xlsx] = list(folder.glob("*.xlsx"))
    wb = load_workbook(xlsx, read_only=True)
    rows = [list(r) for r in wb.active.iter_rows(values_only=True)]
    wb.close()
    header_index = next(i for i, r in enumerate(rows) if r and r[0] == "Transaction Date")
    header = [RENAMES.get(h, h) for h in rows[header_index] if h]
    out = folder / "gateway_export.csv"
    with out.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["Acme Gateway settlement detail"])
        w.writerow(header)
        for r in rows[header_index + 1 :]:
            r = list(r[: len(header)])
            for i in (0, 1):
                if hasattr(r[i], "strftime"):
                    r[i] = f"{r[i].month}/{r[i].day}/{r[i].year}"
            w.writerow(r)
    xlsx.unlink()


@pytest.fixture
def real_project(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    for key, value in make_env(tmp_path, mode="demo").items():
        monkeypatch.setenv(key, value)
    monkeypatch.delenv(MODE_ENV_VAR)  # the mode must come from settings.yaml
    return tmp_path


def test_real_mode_end_to_end(real_project: Path, capsys: pytest.CaptureFixture) -> None:
    root = real_project
    # 2. Portfolio config: add the real portfolio next to the demo one.
    portfolios = yaml.safe_load((root / "config" / "portfolios.yaml").read_text())
    portfolios["portfolios"].append(FAKE_REAL_PORTFOLIO.model_dump())
    (root / "config" / "portfolios.yaml").write_text(yaml.safe_dump(portfolios))

    # 5a. Switch mode in settings.yaml only.
    settings_file = root / "config" / "settings.yaml"
    raw_settings = yaml.safe_load(settings_file.read_text())
    raw_settings.update(data_mode="real", as_of_date=AS_OF.isoformat())
    settings_file.write_text(yaml.safe_dump(raw_settings))
    settings = load_settings()
    assert settings.data_mode == "real"
    assert settings.raw_dir == root / "data" / "raw"

    # 3. Drop files into data/raw/.
    truth = write_fake_real_hotel(settings.raw_dir, AS_OF, settings)

    # 4. One export in a new format, plus a new mapping YAML. No code.
    _convert_processor_to_new_format(
        settings.raw_dir / "fixture_hotels" / "FX01" / "processor_settlement"
    )
    (root / "mappings" / "acme_gateway_settlement.yaml").write_text(
        yaml.safe_dump(NEW_GATEWAY_MAPPING)
    )

    # 5b. Import, match, report through the CLI.
    assert cli.main(["import"]) == 0
    assert cli.main(["match"]) == 0
    out_dir = root / "exports"
    assert cli.main(["report", "--portfolio", "fixture_hotels", "--out", str(out_dir)]) == 0
    printed = capsys.readouterr().out
    assert "acme_gateway_settlement" in printed

    with session_scope(settings) as s:
        files = s.scalars(select(SourceFile)).all()
        assert {f.data_mode for f in files} == {"real"}
        assert any(f.mapping_name == "acme_gateway_settlement" for f in files)
        results = {
            e.vcc_key: (e.exception_type, e.match_method) for e in s.scalars(select(Exception_))
        }
    expected = {k: tuple(v) for k, v in truth["properties"]["FX01"]["results"].items()}
    assert results == expected
    assert any(t != "CHARGED_OK" and t != "NOT_YET_DUE" for t, _ in expected.values())

    # Real data never touches the demo database.
    assert (root / "data" / "leakguard_real.db").exists()
    assert not (root / "data" / "leakguard_demo.db").exists()

    xlsx = next(out_dir.glob("*.xlsx"))
    pdf = next(out_dir.glob("*.pdf"))
    wb = load_workbook(xlsx)
    assert wb.sheetnames == [
        "Summary",
        "By property",
        "By management company",
        "Line items",
        "Method",
    ]
    assert wb["Summary"]["A2"].value.startswith("Fixture Hotels LLC")
    assert pdf.read_bytes().startswith(b"%PDF")
    dispose_engine(settings)
