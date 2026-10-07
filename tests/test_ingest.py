"""Readers, mapping application and the card scanner."""

from pathlib import Path

import pytest
from openpyxl import Workbook

from leakguard.config import PROJECT_ROOT
from leakguard.ingest.mapping import (
    apply_mapping,
    choose_mapping,
    load_all_mappings,
    parse_mapping,
)
from leakguard.ingest.readers import UnsupportedFile, read_table

MAPPINGS, _ = load_all_mappings(PROJECT_ROOT / "mappings")


def by_name(name: str):
    return next(m for m in MAPPINGS if m.name == name)


# Readers ----------------------------------------------------------------------


def test_semicolon_cp1252_csv(tmp_path: Path) -> None:
    path = tmp_path / "res.csv"
    path.write_bytes("Confirmation #;Guest Last Name;Amount\n1;Muñoz;10\n".encode("cp1252"))
    table = read_table(path)
    assert table.header == ["Confirmation #", "Guest Last Name", "Amount"]
    assert table.rows == [["1", "Muñoz", "10"]]
    assert table.delimiter == ";"


def test_tab_delimited(tmp_path: Path) -> None:
    path = tmp_path / "x.txt"
    path.write_text("A\tB\tC\n1\t2\t3\n", encoding="utf-8")
    assert read_table(path).header == ["A", "B", "C"]


def test_title_rows_above_header_in_xlsx(tmp_path: Path) -> None:
    path = tmp_path / "s.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["nothing here"])
    data = wb.create_sheet("Transactions")
    data.append(["Merchant settlement report"])
    data.append(["Period: Jan"])
    data.append([])
    data.append(["Transaction Date", "Card Number", "Amount"])
    data.append(["2026-01-05", "************1234", 12.5])
    wb.save(path)
    table = read_table(
        path, known_headers={"transactiondate", "cardnumber", "amount"}, sheet="Transactions"
    )
    assert table.header == ["Transaction Date", "Card Number", "Amount"]
    assert table.rows[0] == ["2026-01-05", "************1234", "12.5"]
    assert table.first_row_number == 5
    assert table.sheet_names == ["Summary", "Transactions"]


def test_missing_sheet_is_a_clear_error(tmp_path: Path) -> None:
    path = tmp_path / "s.xlsx"
    Workbook().save(path)
    with pytest.raises(ValueError, match="not found"):
        read_table(path, sheet="Nope")


def test_pdf_is_unsupported_with_advice(tmp_path: Path) -> None:
    path = tmp_path / "report.pdf"
    path.write_bytes(b"%PDF-1.4")
    with pytest.raises(UnsupportedFile, match="CSV or Excel"):
        read_table(path)


# Mapping ----------------------------------------------------------------------


def test_all_shipped_mappings_load() -> None:
    _, errors = load_all_mappings(PROJECT_ROOT / "mappings")
    assert errors == []
    assert len(MAPPINGS) >= 7


def test_header_matching_ignores_case_and_punctuation() -> None:
    m = by_name("expedia_vcc_demo")
    resolved = m.resolve(["reservation-id", "CARD AMOUNT", "card_expiry"])
    assert resolved == {"ota_confirmation_no": 0, "vcc_amount": 1, "expiry_date": 2}


def test_best_fitting_mapping_is_chosen() -> None:
    opera = ["Conf No", "External Ref", "Arrival", "Departure", "Res Status", "Guest Name"]
    generic = ["Confirmation #", "Arrival Date", "Departure Date", "Status", "Guest Last Name"]
    assert choose_mapping(MAPPINGS, "pms_reservations", opera).name == "pms_opera_demo_reservations"
    assert choose_mapping(MAPPINGS, "pms_reservations", generic).name == "pms_generic_reservations"
    assert choose_mapping(MAPPINGS, "pms_reservations", ["Foo", "Bar"]) is None


def test_bad_mapping_is_rejected() -> None:
    with pytest.raises(ValueError, match="not a field"):
        parse_mapping(
            {
                "name": "x",
                "source_name": "expedia_vcc",
                "target": "vcc_records",
                "columns": {"nonsense": ["A"]},
            }
        )


def _vcc_table(tmp_path: Path, header: list[str], rows: list[list[str]]):
    path = tmp_path / "cards.csv"
    lines = [",".join(header)] + [",".join(r) for r in rows]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return read_table(path)


def test_full_card_number_in_last4_column_keeps_last4_and_warns(tmp_path: Path) -> None:
    table = _vcc_table(
        tmp_path,
        ["Reservation ID", "Card Number", "Card Amount", "Card Expiry"],
        [["7001", "5555555555554444", "100.00", "2026-12-01"]],
    )
    result = apply_mapping(by_name("expedia_vcc_demo"), table, "cards.csv")
    assert result.rows[0]["card_last4"] == "4444"
    assert "5555555555554444" not in str(result.rows)
    assert result.warnings[0].startswith("Full card numbers found in cards.csv. Dropped.")


def test_unmapped_column_with_card_numbers_is_dropped(tmp_path: Path) -> None:
    table = _vcc_table(
        tmp_path,
        ["Reservation ID", "Card Last 4", "Card Amount", "Card Expiry", "Notes", "CVV"],
        [["7001", "4444", "100.00", "2026-12-01", "paid with 4111111111111111", "123"]],
    )
    result = apply_mapping(by_name("expedia_vcc_demo"), table, "cards.csv")
    assert len(result.rows) == 1
    assert "4111111111111111" not in str(result.rows)
    assert "123" not in str(result.rows[0].values())
    assert any("security code" in w for w in result.warnings)
    assert any(w.startswith("Full card numbers found") for w in result.warnings)


def test_card_number_in_mapped_non_card_column_rejects_row(tmp_path: Path) -> None:
    table = _vcc_table(
        tmp_path,
        ["Reservation ID", "Card Last 4", "Card Amount", "Card Expiry", "Card Status"],
        [
            ["7001", "4444", "100.00", "2026-12-01", "4111111111111111"],
            ["7002", "1234", "50.00", "2026-12-01", "Active"],
        ],
    )
    result = apply_mapping(by_name("expedia_vcc_demo"), table, "cards.csv")
    assert [r["ota_confirmation_no"] for r in result.rows] == ["7002"]
    assert result.rows_rejected == 1
    assert "full card number" in result.issues[0].message
    assert "4111111111111111" not in result.issues[0].message


def test_rejected_rows_report_row_number_and_reason(tmp_path: Path) -> None:
    table = _vcc_table(
        tmp_path,
        ["Reservation ID", "Card Last 4", "Card Amount", "Card Expiry"],
        [["7001", "4444", "100.00", "2026-12-01"], ["", "1111", "oops", "2026-12-01"]],
    )
    result = apply_mapping(by_name("expedia_vcc_demo"), table, "cards.csv")
    assert result.rows_read == 2 and result.rows_rejected == 1
    rows = {i.row_number for i in result.issues}
    assert rows == {3}
    messages = " ".join(i.message for i in result.issues)
    assert "not a money amount" in messages


def test_ambiguous_dates_warn(tmp_path: Path) -> None:
    table = _vcc_table(
        tmp_path,
        ["Reservation ID", "Card Last 4", "Card Amount", "Card Expiry"],
        [["7001", "4444", "100.00", "03/04/2026"]],
    )
    result = apply_mapping(by_name("expedia_vcc_demo"), table, "cards.csv")
    assert any("month-first or day-first" in i.message for i in result.issues)
    assert str(result.rows[0]["expiry_date"]) == "2026-03-04"


def test_reversal_flag_makes_amount_negative(tmp_path: Path) -> None:
    path = tmp_path / "pay.csv"
    path.write_text(
        "Confirmation #,Date,Payment Method,Card Last Four,Amount,Reversal\n"
        "100,01/05/2026,Visa,1234,$50.00,Y\n100,01/05/2026,Visa,1234,$50.00,N\n",
        encoding="utf-8",
    )
    result = apply_mapping(by_name("pms_generic_payments"), read_table(path), "pay.csv")
    assert [r["amount_cents"] for r in result.rows] == [-5000, 5000]
    assert result.rows[0]["is_reversal"] is True


def test_unknown_enum_value_falls_back_with_one_warning(tmp_path: Path) -> None:
    path = tmp_path / "res.csv"
    path.write_text(
        "Confirmation #,Channel,Arrival Date,Departure Date,Status\n"
        "1,Carrier Pigeon,01/15/2026,01/16/2026,Checked Out\n"
        "2,Carrier Pigeon,01/15/2026,01/16/2026,Checked Out\n",
        encoding="utf-8",
    )
    result = apply_mapping(by_name("pms_generic_reservations"), read_table(path), "res.csv")
    assert [r["channel"] for r in result.rows] == ["other", "other"]
    warnings = [i for i in result.issues if i.severity == "warning"]
    assert len(warnings) == 1 and "(2 rows)" in warnings[0].message


def test_guest_names_can_be_hashed(tmp_path: Path) -> None:
    table = _vcc_table(
        tmp_path,
        ["Reservation ID", "Card Last 4", "Card Amount", "Card Expiry", "Guest Name"],
        [["7001", "4444", "100.00", "2026-12-01", "Ann Smith"]],
    )
    plain = apply_mapping(by_name("expedia_vcc_demo"), table, "c.csv")
    hashed = apply_mapping(by_name("expedia_vcc_demo"), table, "c.csv", hash_names=True)
    assert plain.rows[0]["guest_last_name"] == "Smith"
    assert hashed.rows[0]["guest_last_name"] != "Smith"
    assert len(hashed.rows[0]["guest_last_name"]) == 16
